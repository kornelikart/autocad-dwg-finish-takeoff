"""Выгрузка `ProjectResult` в Excel: пять листов, числа в ячейках — числами, округлённые
по общим правилам (м² и прочее — до 0.1, м³ — до 0.01, «от половины вверх»)."""

import re
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from finishcalc.calc import ZoneArea
from finishcalc.model import FloorLayer, FloorType, MaterialLine, Opening, Room
from finishcalc.pipeline import ProjectResult

_MAX_COLUMN_WIDTH = 60
_WIDTH_PADDING = 2
_DEFAULT_DIGITS = 1
_DIGITS = {"м³": 2}
_UNSPECIFIED_MARK = " (не указано)"
_NO_NORM_BASIS = "чертёж (норма не задана)"
_UNKNOWN_FLOOR_TYPE = "тип не описан в экспликации полов"
_DEFAULT_SIZE_MARK = " (размер по умолчанию)"
_OPENING_KIND = {"door": "дверь", "window": "окно"}

_FINISH_SHEET = "Ведомость отделки"
_FLOORS_SHEET = "Экспликация полов"
_MATERIALS_SHEET = "Материалы"
_ROOMS_SHEET = "По помещениям"
_ISSUES_SHEET = "Проверки"

_FINISH_HEADER = ["№ п/п", "Наименование помещения", "Потолок", "Площадь потолка, м²", "Стены/перегородки",
                  "Площадь стен, м²", "Напольное покрытие", "Площадь пола, м²", "Зоны стен", "Примечание"]
_FLOORS_HEADER = ["Тип пола", "Помещения", "Площадь, м²", "Состав"]
_MATERIALS_HEADER = ["Материал", "Ед.", "Количество", "Запас, %", "Основание"]
_ROOMS_HEADER = ["№", "Наименование", "Этаж", "Тип пола", "Площадь пола, м²", "Периметр, м", "Высота, м",
                 "Площадь стен, м²", "Площадь потолка, м²", "Заход на стену, м²", "Проёмы"]
_ISSUES_HEADER = ["Код", "Помещение", "Сообщение"]


def export_excel(result: ProjectResult, path: Path) -> Path:
    """Записывает результат в .xlsx с листами «Ведомость отделки», «Экспликация полов», «Материалы»,
    «По помещениям», «Проверки»; возвращает `path`."""
    path = Path(path)
    workbook = Workbook()
    workbook.remove(workbook.active)
    _write_sheet(workbook, _FINISH_SHEET, _FINISH_HEADER, _finish_rows(result))
    _write_sheet(workbook, _FLOORS_SHEET, _FLOORS_HEADER, _floor_rows(result))
    _write_sheet(workbook, _MATERIALS_SHEET, _MATERIALS_HEADER, [_material_row(m) for m in result.materials])
    _write_sheet(workbook, _ROOMS_SHEET, _ROOMS_HEADER, _room_rows(result))
    _write_sheet(workbook, _ISSUES_SHEET, _ISSUES_HEADER, [[i.code, i.room or "", i.message] for i in result.issues])
    workbook.save(path)
    return path


# --- листы ---

def _finish_rows(result: ProjectResult) -> list[list]:
    """Помещения с отделкой по возрастанию номера. Тексты берутся из описания отделки,
    а если его нет — из строки ведомости (строки ведомости без контура сюда не попадают)."""
    rows = []
    for room in sorted((r for r in result.rooms if r.finish is not None), key=_number_key):
        areas = result.areas[room.number]
        schedule = result.finish_rows.get(room.number)

        def text(key: str) -> str:
            return room.finish.source.get(key) or getattr(schedule, key, "")

        rows.append([
            room.number, room.name or getattr(schedule, "name", ""),
            text("ceiling"), _round(areas.ceiling_m2),
            text("walls"), _round(areas.walls_m2),
            text("floor"), _round(areas.floor_m2),
            _zones_text(areas.zones), getattr(schedule, "note", ""),
        ])
    return rows


def _floor_rows(result: ProjectResult) -> list[list]:
    """Типы пола экспликации в её порядке, затем типы с плана, которых в экспликации нет.
    Площадь и номера помещений — с плана; если на плане типа нет, вместо номеров — подпись экспликации."""
    numbers: dict[str, list[str]] = {}
    for room in sorted(result.rooms, key=_number_key):
        if room.floor_type:
            numbers.setdefault(room.floor_type, []).append(room.number)
    codes = [*result.floor_types, *(c for c in result.floor_type_areas if c not in result.floor_types)]
    rows = []
    for code in codes:
        floor_type = result.floor_types.get(code)
        rooms_text = ", ".join(numbers.get(code, [])) or (floor_type.rooms_label if floor_type else "")
        rows.append([code, rooms_text, _round(result.floor_type_areas.get(code, 0.0)), _composition(floor_type)])
    return rows


def _material_row(line: MaterialLine) -> list:
    basis = _NO_NORM_BASIS if line.basis == "чертёж" else line.basis
    return [line.material, line.unit, _round(line.quantity, _DIGITS.get(line.unit, _DEFAULT_DIGITS)),
            _round(line.reserve_pct), basis]


def _room_rows(result: ProjectResult) -> list[list]:
    rows = []
    for room in sorted(result.rooms, key=_number_key):
        areas = result.areas[room.number]
        schedule = result.finish_rows.get(room.number)
        rows.append([
            room.number, room.name or getattr(schedule, "name", ""), room.floor or "", room.floor_type or "",
            _round(areas.floor_m2), _round(room.perimeter_m), _round(areas.height_m),
            _round(areas.walls_m2), _round(areas.ceiling_m2), _round(areas.upturn_m2),
            "; ".join(_opening_text(o) for o in room.openings),
        ])
    return rows


# --- запись ---

def _write_sheet(workbook: Workbook, title: str, header: list[str], rows: list[list]) -> None:
    """Лист с жирной закреплённой шапкой; ширина колонок по содержимому, не больше предела."""
    sheet = workbook.create_sheet(title)
    widths = [0] * len(header)
    for row_index, row in enumerate([header, *rows], start=1):
        for column, value in enumerate(row):
            cell = sheet.cell(row_index, column + 1, None if value == "" else value)
            if cell.data_type == "f":  # текст с чертежа, начинающийся с «=», не должен стать формулой
                cell.data_type = "s"
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if row_index == 1:
                cell.font = Font(bold=True)
            widths[column] = max(widths[column], len(str(value)))
    for column, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = min(width + _WIDTH_PADDING, _MAX_COLUMN_WIDTH)
    sheet.freeze_panes = "A2"


# --- текст и числа ---

def _round(value: float, digits: int = _DEFAULT_DIGITS) -> float:
    """Округление «от половины вверх»; `round(…, 9)` убирает шум плавающей запятой."""
    return float(Decimal(repr(round(value, 9))).quantize(Decimal(1).scaleb(-digits), rounding=ROUND_HALF_UP))


def _comma(value: float) -> str:
    return f"{value:.1f}".replace(".", ",")


def _num(value: float) -> str:
    """Размер без лишних нулей, десятичная запятая: 0.9 -> «0,9», 1.0 -> «1»."""
    return f"{round(value, 2):g}".replace(".", ",")


def _height(value: float) -> str:
    """Граница полосы стены: «0», десятые («1,8», «3,0»), сотые только если без них терялась бы точность."""
    if value == 0:
        return "0"
    text = f"{value:.2f}"
    return (text[:-1] if text.endswith("0") else text).replace(".", ",")


def _zones_text(zones: list[ZoneArea]) -> str:
    parts = []
    for zone in zones:
        part = f"{_height(zone.from_m)}–{_height(zone.to_m)} м: {_comma(_round(zone.area_m2))} м²"
        parts.append(part + _UNSPECIFIED_MARK if zone.unspecified else part)
    return "; ".join(parts)


def _composition(floor_type: FloorType | None) -> str:
    if floor_type is None:
        return _UNKNOWN_FLOOR_TYPE
    if not floor_type.layers:
        return floor_type.source_text
    return "; ".join(_layer_text(layer) for layer in floor_type.layers)


def _layer_text(layer: FloorLayer) -> str:
    known = [t for t in (layer.thickness_mm_min, layer.thickness_mm_max) if t is not None]
    if not known:
        return layer.material
    low, high = min(known), max(known)
    thickness = _num(low) if low == high else f"{_num(low)}–{_num(high)}"
    return f"{layer.material}, {thickness} мм"


def _opening_text(opening: Opening) -> str:
    parts = [_OPENING_KIND[opening.kind]]
    if opening.mark:
        parts.append(opening.mark)
    parts.append(f"{_num(opening.width_m)}×{_num(opening.height_m)} м")
    text = " ".join(parts)
    return text + _DEFAULT_SIZE_MARK if opening.size_source == "default" else text


def _number_key(room: Room) -> list[tuple[int, int, str]]:
    """Естественный порядок номеров: Т2 перед Т10, цифры перед буквами."""
    return [(0, int(part), "") if part.isdigit() else (1, 0, part)
            for part in re.split(r"(\d+)", room.number) if part]
