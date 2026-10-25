from dataclasses import replace

import openpyxl
import pytest

from finishcalc.export import export_excel
from finishcalc.extract.tables import FinishRow
from finishcalc.model import FloorLayer, FloorType, Issue, MaterialLine, Opening

SHEETS = ["Ведомость отделки", "Экспликация полов", "Материалы", "По помещениям", "Проверки"]
ZONES_TEXT = "0–1,8 м: 24,3 м²; 1,8–2,1 м: 4,7 м² (не указано); 2,1–3,0 м: 15,8 м²"


def load(result, tmp_path):
    return openpyxl.load_workbook(export_excel(result, tmp_path / "out.xlsx"))


def values(ws):
    return [[c.value for c in row] for row in ws.iter_rows()]


def test_export(tmp_path, small_result):
    p = export_excel(small_result, tmp_path / "out.xlsx")
    wb = openpyxl.load_workbook(p)
    assert wb.sheetnames == ["Ведомость отделки", "Экспликация полов", "Материалы", "По помещениям", "Проверки"]
    ws = wb["Ведомость отделки"]
    assert [c.value for c in ws[1]][:4] == ["№ п/п", "Наименование помещения", "Потолок", "Площадь потолка, м²"]
    assert ws.cell(2, 1).value == "Т1" and ws.cell(2, 6).value == 44.7
    assert "1,8–2,1 м: 4,7 м² (не указано)" in ws.cell(2, 9).value


def test_export_returns_the_path(tmp_path, small_result):
    path = tmp_path / "out.xlsx"
    assert export_excel(small_result, path) == path and path.exists()


def test_finish_sheet_header_and_row(tmp_path, small_result):
    ws = load(small_result, tmp_path)["Ведомость отделки"]
    assert values(ws) == [
        ["№ п/п", "Наименование помещения", "Потолок", "Площадь потолка, м²", "Стены/перегородки",
         "Площадь стен, м²", "Напольное покрытие", "Площадь пола, м²", "Зоны стен", "Примечание"],
        ["Т1", "Санузел", "Подвесной потолок", 20.0, "Плитка до 1800 мм, окраска выше 2100 мм", 44.7,
         "Керамогранит, завести на стену на 300мм.", 20.0, ZONES_TEXT, "Вход с коридора"],
    ]


def test_areas_are_numbers_not_text(tmp_path, small_result):
    wb = load(small_result, tmp_path)
    numeric = {"Ведомость отделки": (4, 6, 8), "Экспликация полов": (3,), "Материалы": (3, 4),
               "По помещениям": (5, 6, 7, 8, 9, 10)}
    for title, columns in numeric.items():
        for column in columns:
            cell = wb[title].cell(2, column)
            assert cell.data_type == "n" and isinstance(cell.value, (int, float)), (title, column)


def test_areas_round_half_up(tmp_path, small_result):
    areas = {"Т1": replace(small_result.areas["Т1"], floor_m2=20.25, ceiling_m2=20.35)}
    ws = load(replace(small_result, areas=areas), tmp_path)["Ведомость отделки"]
    assert (ws.cell(2, 8).value, ws.cell(2, 4).value) == (20.3, 20.4)  # round() дал бы 20.2


def test_finish_sheet_lists_only_rooms_with_finish_in_ascending_order(tmp_path, small_result):
    room = small_result.rooms[0]
    rooms = [replace(room, number="Т10"), replace(room, number="Т2"), room, replace(room, number="Т3", finish=None)]
    areas = {r.number: small_result.areas["Т1"] for r in rooms}
    ws = load(replace(small_result, rooms=rooms, areas=areas), tmp_path)["Ведомость отделки"]
    assert [row[0] for row in values(ws)[1:]] == ["Т1", "Т2", "Т10"]


def test_finish_row_without_contour_is_not_listed_and_does_not_break_export(tmp_path, small_result):
    rows = {**small_result.finish_rows,
            "А07": FinishRow("А07", "Кладовая", "Окраска", "Окраска", "Линолеум", None, "")}
    ws = load(replace(small_result, finish_rows=rows), tmp_path)["Ведомость отделки"]
    assert [row[0] for row in values(ws)[1:]] == ["Т1"]


def test_empty_finish_source_falls_back_to_schedule_row(tmp_path, small_result):
    room = small_result.rooms[0]
    room.finish = replace(room.finish, source={})
    room.name = ""
    ws = load(small_result, tmp_path)["Ведомость отделки"]
    row = small_result.finish_rows["Т1"]
    assert [ws.cell(2, c).value for c in (2, 3, 5, 7)] == [row.name, row.ceiling, row.walls, row.floor]


def test_finish_source_wins_over_schedule_row(tmp_path, small_result):
    rows = {"Т1": replace(small_result.finish_rows["Т1"], ceiling="Окраска")}
    ws = load(replace(small_result, finish_rows=rows), tmp_path)["Ведомость отделки"]
    assert ws.cell(2, 3).value == "Подвесной потолок"


def test_room_without_schedule_row_has_empty_note(tmp_path, small_result):
    ws = load(replace(small_result, finish_rows={}, mode="B"), tmp_path)["Ведомость отделки"]
    assert ws.cell(2, 10).value is None and ws.cell(2, 9).value == ZONES_TEXT


def test_zone_bounds_keep_tenths_and_show_hundredths_when_needed(tmp_path, small_result):
    areas = small_result.areas["Т1"]
    zones = [replace(areas.zones[0], from_m=0.0, to_m=1.75), replace(areas.zones[1], from_m=1.75, to_m=2.4)]
    ws = load(replace(small_result, areas={"Т1": replace(areas, zones=zones)}), tmp_path)["Ведомость отделки"]
    assert ws.cell(2, 9).value.startswith("0–1,75 м: 24,3 м²; 1,75–2,4 м: 4,7 м² (не указано)")


def test_floor_explication_sheet(tmp_path, small_result):
    ws = load(small_result, tmp_path)["Экспликация полов"]
    assert values(ws) == [["Тип пола", "Помещения", "Площадь, м²", "Состав"],
                          ["7", "Т1", 20.0, "Керамогранит, 20 мм; Стяжка, 50 мм"]]


def test_floor_explication_lists_plan_rooms_and_unknown_types(tmp_path, small_result):
    room = small_result.rooms[0]
    other = replace(room, number="Т2", floor_type="9")
    types = {**small_result.floor_types,
             "5": FloorType("5", "Кабинеты", [], "Линолеум 3 мм"),
             "6": FloorType("6", "", [FloorLayer("Ламинат", 7, 9, "area", "laminate"),
                                      FloorLayer("Подложка", None, None, "piece", "other")], "")}
    result = replace(small_result, rooms=[other, room, replace(room, number="Т10")],
                     areas={n: small_result.areas["Т1"] for n in ("Т1", "Т2", "Т10")},
                     floor_types=types, floor_type_areas={"7": 40.0, "9": 20.0})
    rows = {row[0]: row for row in values(load(result, tmp_path)["Экспликация полов"])[1:]}
    assert rows["7"][1:3] == ["Т1, Т10", 40.0]
    assert rows["5"][1:] == ["Кабинеты", 0.0, "Линолеум 3 мм"]  # на плане нет: подпись из экспликации
    assert rows["6"][3] == "Ламинат, 7–9 мм; Подложка"
    assert rows["9"][1:3] == ["Т2", 20.0] and "не описан" in rows["9"][3]
    assert list(rows) == ["7", "5", "6", "9"]


def test_materials_sheet(tmp_path, small_result):
    ws = load(small_result, tmp_path)["Материалы"]
    rows = values(ws)
    assert rows[0] == ["Материал", "Ед.", "Количество", "Запас, %", "Основание"]
    by_name = {row[0]: row for row in rows[1:]}
    assert by_name["Подвесной потолок"] == ["Подвесной потолок", "м²", 21.0, 5.0, "норма"]
    assert by_name["Клей плиточный"][1:3] == ["кг", 121.5]
    assert by_name["Стяжка"][1:3] == ["м³", 1.0]
    assert len(rows) - 1 == len(small_result.materials)


def test_materials_without_norm_are_marked(tmp_path, small_result):
    lines = [MaterialLine("декоративная панель", "м²", 49.94, 0.0, "чертёж", ["Т1"]),
             MaterialLine("Бетон", "м³", 1.005, 5.0, "норма", ["Т1"])]
    rows = values(load(replace(small_result, materials=lines), tmp_path)["Материалы"])[1:]
    assert rows == [["декоративная панель", "м²", 49.9, 0.0, "чертёж (норма не задана)"],
                    ["Бетон", "м³", 1.01, 5.0, "норма"]]


def test_rooms_sheet(tmp_path, small_result):
    ws = load(small_result, tmp_path)["По помещениям"]
    assert values(ws) == [
        ["№", "Наименование", "Этаж", "Тип пола", "Площадь пола, м²", "Периметр, м", "Высота, м",
         "Площадь стен, м²", "Площадь потолка, м²", "Заход на стену, м²", "Проёмы"],
        ["Т1", "Санузел", "1", "7", 20.0, 18.0, 3.0, 44.7, 20.0, 5.1,
         "дверь ДС-2 0,9×2,1 м; окно 3 1,5×1,5 м (размер по умолчанию)"],
    ]


def test_rooms_sheet_lists_every_room_sorted_with_blank_unknowns(tmp_path, small_result):
    room = small_result.rooms[0]
    bare = replace(room, number="Т2", name="", floor=None, floor_type=None, finish=None,
                   openings=[Opening("door", None, 1.0, 2.0, 0.0, ["Т2"], "default", (0, 0))])
    areas = {"Т1": small_result.areas["Т1"], "Т2": small_result.areas["Т1"]}
    ws = load(replace(small_result, rooms=[bare, room], areas=areas), tmp_path)["По помещениям"]
    rows = values(ws)[1:]
    assert [r[0] for r in rows] == ["Т1", "Т2"]
    assert rows[1][1:4] == [None, None, None]
    assert rows[1][10] == "дверь 1×2 м (размер по умолчанию)"


def test_checks_sheet(tmp_path, small_result):
    ws = load(small_result, tmp_path)["Проверки"]
    rows = values(ws)
    assert rows[0] == ["Код", "Помещение", "Сообщение"]
    assert rows[1][:2] == ["unspecified_wall_band", "Т1"]
    assert "4,7" in rows[1][2]
    assert rows[2] == ["reveals_not_counted", None, "Откосы проёмов не учтены"]


def test_any_issue_code_is_written_as_is(tmp_path, small_result):
    issues = [Issue("duplicate_room_number", "А01", "Номер А01 стоит в нескольких контурах")]
    ws = load(replace(small_result, issues=issues), tmp_path)["Проверки"]
    assert values(ws)[1] == ["duplicate_room_number", "А01", "Номер А01 стоит в нескольких контурах"]


def test_text_from_drawings_is_never_a_formula(tmp_path, small_result):
    issues = [Issue("opening_without_mark", "Т1", '=HYPERLINK("http://example.invalid","x")')]
    cell = load(replace(small_result, issues=issues), tmp_path)["Проверки"].cell(2, 3)
    assert cell.data_type == "s" and cell.value == '=HYPERLINK("http://example.invalid","x")'


def test_empty_result_gives_five_sheets_with_headers_only(tmp_path, small_result):
    empty = replace(small_result, rooms=[], areas={}, finish_rows={}, floor_types={}, floor_type_areas={},
                    materials=[], issues=[])
    wb = load(empty, tmp_path)
    assert wb.sheetnames == SHEETS
    assert all(ws.max_row == 1 for ws in wb)


def test_headers_are_bold_first_row_is_frozen_widths_are_capped(tmp_path, small_result):
    long_text = "Очень длинное предупреждение " * 10
    issues = [Issue("opening_without_mark", "Т1", long_text)]
    wb = load(replace(small_result, issues=issues), tmp_path)
    for ws in wb:
        assert all(c.font.bold for c in ws[1]), ws.title
        assert ws.freeze_panes == "A2", ws.title
        assert all(d.width <= 60 for d in ws.column_dimensions.values()), ws.title
    checks = wb["Проверки"]
    assert checks.column_dimensions["C"].width == 60  # длинный текст упирается в предел
    assert checks.column_dimensions["A"].width == pytest.approx(len("opening_without_mark") + 2)
    assert checks["C2"].alignment.wrap_text
